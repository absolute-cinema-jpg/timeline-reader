"""Lightweight sanity tests. Run with:  ./.venv/bin/python -m pytest -q
(or plain:  ./.venv/bin/python tests/test_core.py)
"""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import tempfile

from timeline_reader.captions import parse_caption_text, to_srt
from timeline_reader.columns import (
    CLIPLIST_REPORT,
    TAGFINDER_REPORT,
    ColumnSelection,
    META_PREFIX,
)
from timeline_reader.effects import classify, is_optical_category
from timeline_reader.exporters import export_table, format_labels
from timeline_reader.models import Clip, Timeline
from timeline_reader.timecode import Timecode


def test_timecode_roundtrip_25():
    tc = Timecode.from_string("11:45:50:01", 25.0)
    assert tc.frames == 1058751
    assert tc.to_string() == "11:45:50:01"


def test_timecode_srt_ms():
    assert Timecode(12, 25.0).to_srt() == "00:00:00,480"
    assert Timecode.from_string("10:00:05:00", 25.0).to_srt() == "10:00:05,000"


def test_effect_classification():
    assert classify("EFF2_BLEND_RESIZE", None) == ("Resize", True)
    assert classify("EFF2_SBLEND", "3DWarp") == ("3D Warp", True)
    assert classify("INT_EFF_SPATIAL_ADAPTER", None)[1] is False
    assert is_optical_category("Timewarp") is True
    assert is_optical_category("Colour Correction") is False


class _Obj:
    """Duck-typed stand-in for a pyavb node (only ``property_data`` matters)."""

    def __init__(self, **data):
        self.property_data = data


def _param(uuid, offsets_values):
    points = [_Obj(offset=off, value=val) for off, val in offsets_values]
    return _Obj(uuid=uuid, control_track=_Obj(control_points=points))


# Well-known Avid parameter UUIDs (from avb.parameter_uuids).
_POS_Y = "8d568126-847e-11d5-935a-50f857c10000"
_SCALE_X = "8d568129-847e-11d5-935a-50f857c10000"
_SCALE_Y = "8d56812a-847e-11d5-935a-50f857c10000"
_CROP_L = "8d56812e-847e-11d5-935a-50f857c10000"
_OFFSET_MAP = "8d56827a-847e-11d5-935a-50f857c10000"
_SPEED_MAP = "8d56827c-847e-11d5-935a-50f857c10000"
_SPEED_OFFSET_MAP = "8d56827d-847e-11d5-935a-50f857c10000"


def test_keyframes_resolve_to_source_timecode():
    from timeline_reader.keyframes import summarize_keyframes

    node = _Obj(length=50, param_list=[
        _param(_POS_Y, [([0, 1], 49.0), ([50, 1], 78.0)]),
        _param(_SCALE_X, [([0, 1], 116.0), ([50, 1], 121.0)]),
        _param(_CROP_L, [([0, 1], 0.0), ([50, 1], 0.0)]),  # constant -> skipped
    ])
    # src_start = 1h in frames -> keyframe at offset 0 == src-in, offset 50 == +2s.
    src_in = 25 * 60 * 60  # 01:00:00:00 at 25 fps
    out = summarize_keyframes(node, src_start=src_in, fps=25.0, drop=False)
    assert "Pos Y: 49@01:00:00:00, 78@01:00:02:00" in out
    assert "Scale X: 116@01:00:00:00, 121@01:00:02:00" in out
    assert "Crop Left" not in out  # unchanged parameter is not a keyframe


def test_keyframes_earlier_offset_gives_earlier_source_tc():
    from timeline_reader.keyframes import summarize_keyframes

    # An elastic keyframe before the clip's in-point must report the *earlier*
    # source timecode, not be clamped to the clip in-point.
    src_in = 25 * 60 * 60  # 01:00:00:00
    node = _Obj(length=100, param_list=[
        _param(_SCALE_X, [([-702, 1], 100.0), ([722, 1], 145.0)]),
    ])
    out = summarize_keyframes(node, src_start=src_in, fps=25.0, drop=False)
    # -702 frames -> 00:59:31:23 ; +722 frames -> 01:00:28:22
    assert out == "Scale X: 100@00:59:31:23, 145@01:00:28:22"


def test_keyframes_long_curve_is_thinned_to_endpoints():
    from timeline_reader.keyframes import summarize_keyframes

    pts = [([i, 1], float(i)) for i in range(20)]  # 20 points > cap
    node = _Obj(length=19, param_list=[_param(_SCALE_X, pts)])
    out = summarize_keyframes(node, src_start=0, fps=25.0, drop=False)
    assert "(20 kfs)" in out
    assert out.count("@") == 2  # only first and last shown


def test_keyframes_pairs_x_and_y():
    from timeline_reader.keyframes import summarize_keyframes

    node = _Obj(length=50, param_list=[
        _param(_SCALE_X, [([0, 1], 100.0), ([50, 1], 145.0)]),
        _param(_SCALE_Y, [([0, 1], 100.0), ([50, 1], 145.0)]),
    ])
    out = summarize_keyframes(node, src_start=0, fps=25.0, drop=False)
    # X/Y are combined into one (x, y) pair per keyframe, not listed separately.
    assert out == "Scale: (100, 100)@00:00:00:00, (145, 145)@00:00:02:00"


def test_keyframes_static_crop_is_reported():
    from timeline_reader.keyframes import summarize_keyframes

    # A plain crop: no keyframes, but the non-zero crop amounts are still shown.
    node = _Obj(length=124, param_list=[
        _Obj(uuid=_CROP_L, value=394.0, control_track=None),
    ])
    out = summarize_keyframes(node, src_start=0, fps=25.0, drop=False)
    assert out == "Crop L 394.0"  # crop reported to 1 dp


def _mparam(uuid, offsets_values):
    return _param(uuid, offsets_values)


def test_motion_freeze_frame():
    from timeline_reader.keyframes import describe_motion

    # Source offset never advances -> freeze, despite a bogus [length,1] speed_ratio.
    node = _Obj(length=133, speed_ratio=[133, 1], param_list=[
        _mparam(_SPEED_OFFSET_MAP, [([0, 1], 0.0), ([133, 1], 0.0)]),
        _mparam(_SPEED_MAP, [([0, 1], 0.0)]),
    ])
    assert describe_motion(node, src_start=0, fps=25.0, drop=False) == ("Freeze Frame", "")


def test_motion_trim_to_fill_speed():
    from timeline_reader.keyframes import describe_motion

    # A raw offset map (no speed map) playing 145 source frames over 69 output -> ~210%.
    node = _Obj(length=69, speed_ratio=[69, 146], param_list=[
        _mparam(_OFFSET_MAP, [([0, 1], 0.0), ([68, 1], 145.0)]),
    ])
    cat, detail = describe_motion(node, src_start=0, fps=25.0, drop=False)
    assert cat == "Trim to Fill"
    assert detail == "210.14%"  # 145 src / 69 output, to 2 dp


def test_motion_variable_speed_reports_percent_at_source_tc():
    from timeline_reader.keyframes import describe_motion

    offmap = [([0, 1], 0.0), ([12, 1], 8.0), ([88, 1], 44.0)]
    node = _Obj(length=88, param_list=[
        _mparam(_SPEED_OFFSET_MAP, offmap),
        _mparam(_SPEED_MAP, [([0, 1], 1.0), ([12, 1], 0.35), ([88, 1], 1.0)]),
    ])
    src_in = 25 * 60 * 60  # 01:00:00:00
    cat, detail = describe_motion(node, src_start=src_in, fps=25.0, drop=False)
    assert cat == "Timewarp"
    # speed% at each keyframe, timed by its source offset from the offset map.
    assert detail == "100%@01:00:00:00, 35%@01:00:00:08, 100%@01:00:01:19"


def test_transitions_are_standalone_rows_in_opticals():
    import os

    from timeline_reader.parsers.avb_parser import parse
    from timeline_reader.effects import is_optical_category
    from timeline_reader.timecode import frames_to_tc

    path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "01-test files", "bin", "timeline-reader test.avb")
    if not os.path.exists(path):
        return
    tl = parse(path, key=5)  # the "opticals" sequence

    dissolves = [c for c in tl.clips if any(e.category == "Dissolve" for e in c.effects)]
    assert dissolves, "dissolves should appear as their own optical rows"
    for c in dissolves:
        assert c.is_transition
        assert c.clip_name == "" and c.tape_name == ""      # no clip identity
        assert c.rec_end - c.rec_start == 25                 # its own record span
    # The first dissolve runs to 00:00:29:22 (rec-out), not the incoming clip's out.
    assert frames_to_tc(min(dissolves, key=lambda c: c.rec_start).rec_end, tl.fps) == "00:00:29:22"

    # A fit/trim-to-fill consumes more source than its record length (sped up).
    fill = [c for c in tl.clips if any(e.category == "Trim to Fill" for e in c.effects)]
    assert fill, "trim-to-fill should be detected"
    f = fill[0]
    assert (f.src_end - f.src_start) > (f.rec_end - f.rec_start)


def test_caption_to_srt():
    raw = "<begin subtitles>\n10:00:01:00 10:00:04:12 Hello\nWorld\n<end subtitles>\n"
    doc = parse_caption_text(raw, fps=25.0)
    assert len(doc.cues) == 1
    srt = to_srt(doc)
    assert "10:00:01,000 --> 10:00:04,480" in srt
    assert "Hello\nWorld" in srt


def test_caption_decode_encodings():
    # Avid Caption exports may be UTF-8 or UTF-16 (BOM or bare); all must decode.
    from timeline_reader.captions import _decode
    text = "<begin subtitles>\r\n00:00:01:00 00:00:02:00\r\nHi there\r\n"
    for enc in ("utf-8", "utf-8-sig", "utf-16", "utf-16-le", "utf-16-be"):
        doc = parse_caption_text(_decode(text.encode(enc)), fps=25.0)
        assert len(doc.cues) == 1, f"{enc} decoded to {len(doc.cues)} cues"
        assert doc.cues[0].lines == ["Hi there"]


def _demo_timeline() -> Timeline:
    tl = Timeline(name="demo", fps=25.0)
    tl.add(Clip(index=1, track="V1", clip_name="1-1-3", tape_name="A001",
                src_start=100, src_end=150, rec_start=0, rec_end=50,
                meta={"Take": "3", "Scene": "1"}))
    tl.add(Clip(index=2, track="V1", clip_name="2-2-1", tape_name="A002",
                src_start=200, src_end=240, rec_start=50, rec_end=90,
                meta={"Take": "1"}))
    return tl


def test_default_columns_order():
    tl = _demo_timeline()
    sel = ColumnSelection(CLIPLIST_REPORT.key)  # no overrides -> defaults
    headers, rows = CLIPLIST_REPORT.build(tl, sel)
    # # leads, Clip Name second; the rest follow.
    assert headers == ["#", "Clip Name", "Track", "Tape / Source",
                       "Src In", "Src Out", "Rec In", "Rec Out", "Duration"]
    assert len(rows) == 2


def test_metadata_column_opt_in():
    tl = _demo_timeline()
    assert "Take" in tl.meta_columns and "Scene" in tl.meta_columns
    sel = ColumnSelection(CLIPLIST_REPORT.key)
    sel.set(META_PREFIX + "Take", True)
    headers, rows = CLIPLIST_REPORT.build(tl, sel)
    assert headers[-1] == "Take"
    assert rows[0][-1] == "3" and rows[1][-1] == "1"


def test_tagfinder_report_only_yields_tagged_clips():
    tl = Timeline(name="demo", fps=25.0)
    tl.add(Clip(index=1, track="V1", clip_name="A", note="stock",
                rec_start=0, rec_end=50))
    tl.add(Clip(index=2, track="V1", clip_name="B",
                rec_start=50, rec_end=90))  # no note -> excluded
    tl.add(Clip(index=3, track="V1", clip_name="C", note="Stock footage",
                rec_start=90, rec_end=120))
    sel = ColumnSelection(TAGFINDER_REPORT.key)
    headers, rows = TAGFINDER_REPORT.build(tl, sel)
    assert headers[:2] == ["#", "Note"]
    # Only the two note-bearing clips appear, and the note leads.
    assert [r[1] for r in rows] == ["stock", "Stock footage"]
    assert [r[2] for r in rows] == ["A", "C"]


def test_exclude_default_column():
    tl = _demo_timeline()
    sel = ColumnSelection(CLIPLIST_REPORT.key)
    sel.set("duration", False)
    headers, _ = CLIPLIST_REPORT.build(tl, sel)
    assert "Duration" not in headers


def test_column_reorder():
    tl = _demo_timeline()
    sel = ColumnSelection(CLIPLIST_REPORT.key)
    sel.set(META_PREFIX + "Take", True)
    # Try to move Track and Take ahead of the "#" column; "#" is pinned first
    # regardless, then the listed columns, then the canonical remainder.
    sel.set_order(["track", META_PREFIX + "Take", "index"])
    headers, rows = CLIPLIST_REPORT.build(tl, sel)
    assert headers[:3] == ["#", "Track", "Take"]
    assert headers[3] == "Clip Name"  # unlisted columns stay in canonical order
    assert rows[0][:3] == ["1", "V1", "3"]


def test_marker_extraction():
    """Marker collection resolves absolute position, colour and attributes, and
    de-duplicates locators reached more than once during the track walk.

    Real Avid marker objects aren't in the sample bins, so this exercises the
    extraction with stand-in objects shaped like ``avb.misc.Marker``.
    """
    from timeline_reader.parsers import avb_parser as A

    class FakeMarker:
        def __init__(self, comp_offset, color, attrs):
            self.comp_offset = comp_offset
            self.color = color
            self.attributes = attrs

    class Node:
        property_data: dict = {}

        def __init__(self, attributes):
            self.attributes = attributes

    red = FakeMarker(
        12,
        [90 << 8, 190 << 8, 85 << 8],  # RGB says "Green"...
        {
            "_ATN_CRM_COM": "VFX shot",
            "_ATN_CRM_USER": "alex",
            "_ATN_CRM_COLOR": "Red",  # ...but the named colour is authoritative
            "_ATN_CRM_LENGTH": 5,
        },
    )
    seg = Node({"mark": red})
    tl = Timeline(fps=25.0)
    seen: set = set()
    # Segment starts at record frame 100; marker offset 12 -> absolute 112.
    A._gather_markers(seg, "V1", 100, tl, seen, FakeMarker)
    A._gather_markers(seg, "V1", 100, tl, seen, FakeMarker)  # second pass: no dupes

    assert len(tl.markers) == 1
    m = tl.markers[0]
    assert m.position == 112
    assert m.track == "V1"
    assert m.colour == "Red"  # attribute name wins over the RGB triple
    assert m.comment == "VFX shot"
    assert m.user == "alex"
    assert m.length == 5


def test_marker_extraction_from_sample_bin():
    """End-to-end against the real sample bin's 'markers' sequence, when present.

    The bin lives under the gitignored ``01-test files/``; skip cleanly if absent.
    """
    path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "01-test files", "bin", "timeline-reader test.avb",
    )
    if not os.path.exists(path):
        return
    from timeline_reader.parsers import parse_timeline
    from timeline_reader.parsers.avb_parser import sequences

    key = next(o.key for o in sequences(path) if o.name == "markers")
    tl = parse_timeline(path, key)
    got = {(m.comment, m.colour, m.track) for m in tl.markers}
    assert ("drama 1", "Red", "V1") in got
    assert ("opticals 1", "Cyan", "V3") in got
    assert ("vfx 1", "White", "V4") in got
    assert len(tl.markers) == 7
    assert all(m.user and m.date for m in tl.markers)


def test_timeline_notes_from_sample_bin():
    """Timeline clip notes read from the sample bin's 'timeline notes' sequence,
    each resolving to its clip's full source detail, when the sample is present.
    """
    path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "01-test files", "bin", "timeline-reader test.avb",
    )
    if not os.path.exists(path):
        return
    from timeline_reader.parsers import parse_timeline
    from timeline_reader.parsers.avb_parser import sequences

    key = next(o.key for o in sequences(path) if o.name == "timeline notes")
    tl = parse_timeline(path, key)
    notes = {c.note for c in tl.clips if c.note}
    assert notes == {"Note 1", "note 2", "note 3", "lol"}

    # Only tagged clips reach the report, and a note keeps its clip's metadata.
    ctx = list(TAGFINDER_REPORT.iter_ctx(tl))
    assert len(ctx) == 4
    note1 = next(c for c in ctx if c.clip.note == "Note 1")
    assert note1.clip.tape_name and note1.clip.src_start > 0

    # Case-insensitive substring search is how the tab filters.
    hits = [c for c in ctx if "note" in c.clip.note.lower()]
    assert {c.clip.note for c in hits} == {"Note 1", "note 2", "note 3"}


def test_captions_from_sample_bin():
    """SubCap subtitles read straight from the bin's 'captions no VFX cards'
    sequence — timing and multi-line text — when the gitignored sample is present.
    """
    path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "01-test files", "bin", "timeline-reader test.avb",
    )
    if not os.path.exists(path):
        return
    from timeline_reader.captions import to_srt
    from timeline_reader.captions_avb import caption_sequences, parse_captions
    from timeline_reader.timecode import Timecode

    key = next(o.key for o in caption_sequences(path) if o.name == "captions no VFX cards")
    doc = parse_captions(path, key)
    assert len(doc.cues) == 55
    assert doc.fps == 25.0 and doc.drop is False

    first = doc.cues[0]
    assert Timecode(first.start, 25.0, False).to_string() == "00:02:41:05"
    assert Timecode(first.end, 25.0, False).to_string() == "00:02:43:21"
    assert first.text() == "You wouldn't let go of my hands that day"

    # A wrapped caption keeps both lines (the flat .txt export drops the second).
    multi = next(c for c in doc.cues if c.text().startswith("But you know what"))
    assert multi.text() == "But you know what I’ve been dying to try?\nAiden’s amazing chapati"

    # The SubRip render is well-formed: index, arrow-timed range, then text.
    srt = to_srt(doc)
    assert srt.startswith("1\n00:02:41,200 --> 00:02:43,840\n")


_HEADERS = ["#", "Clip Name", "Rec In"]
_ROWS = [["1", "2-25-3", "00:00:03:17"], ["2", "wide’s", "00:00:05:00"]]


def test_export_formats_listed():
    labels = format_labels()
    assert any(".xlsx" in l for l in labels) and any(".ods" in l for l in labels)


def test_xlsx_roundtrip():
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "r.xlsx")
        export_table(path, _HEADERS, _ROWS, "xlsx", sheet_name="THR:test/1")
        ws = load_workbook(path).active
        assert len(ws.title) <= 31 and "/" not in ws.title  # sheet name sanitised
        assert [c.value for c in ws[1]] == _HEADERS
        assert [c.value for c in ws[2]] == _ROWS[0]
        assert ws["B3"].value == "wide’s"  # unicode preserved


def test_ods_roundtrip():
    from odf import teletype
    from odf.opendocument import load
    from odf.table import Table, TableCell, TableRow
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "r.ods")
        export_table(path, _HEADERS, _ROWS, "ods", sheet_name="Report")
        table = load(path).spreadsheet.getElementsByType(Table)[0]
        rows = table.getElementsByType(TableRow)
        header = [teletype.extractText(c) for c in rows[0].getElementsByType(TableCell)]
        assert header == _HEADERS
        assert teletype.extractText(rows[2].getElementsByType(TableCell)[1]) == "wide’s"


def test_settings_roundtrip():
    from PySide6.QtCore import QSettings
    d = tempfile.mkdtemp()
    QSettings.setDefaultFormat(QSettings.IniFormat)
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, d)  # isolate from real prefs
    from timeline_reader import settings as S
    from timeline_reader.exporters import index_of_kind

    S.set_export_format_kind("xlsx")
    assert S.export_format_kind() == "xlsx"
    assert index_of_kind(S.export_format_kind()) == 2

    S.remember_open_path(os.path.join(d, "bin", "seq.avb"))  # dir must exist to be saved
    os.makedirs(os.path.join(d, "bin"), exist_ok=True)
    S.remember_open_path(os.path.join(d, "bin", "seq.avb"))
    assert S.last_open_dir() == os.path.join(d, "bin")

    S.set_caption_fps_index(4)
    assert S.caption_fps_index() == 4


# --------------------------------------------------------------------------- #
# Music Tracker
# --------------------------------------------------------------------------- #
def _music_clip(track, name, rs, re, *, head=0, tail=0, artist="", title="",
                album="", muted=False):
    meta = {}
    if artist:
        meta["Lead performer(s)/Soloist(s)"] = artist
    if title:
        meta["Title/songname/content descripti"] = title
    if album:
        meta["Album/Movie/Show title"] = album
    return Clip(index=0, track=track, clip_name=name, tape_name=f"{name}.wav",
                rec_start=rs, rec_end=re, fps=25.0,
                head_transition=head, tail_transition=tail, muted=muted, meta=meta)


def _music_timeline(clips, start_tc=0) -> Timeline:
    tl = Timeline(name="demo", fps=25.0)
    tl.audio_clips = clips
    tl.start_tc = start_tc
    return tl


def test_music_merge_gap_and_interrupt():
    from timeline_reader import music
    # 25 fps -> 1s gap = 25 frames.
    tl = _music_timeline([
        _music_clip("A5", "song1", 0, 100),        # opens a cue
        _music_clip("A5", "song1", 110, 200),      # +10f gap, same song -> merges
        _music_clip("A5", "song1", 400, 500),      # +200f gap -> new cue (same song)
        _music_clip("A5", "song2", 505, 600),      # different song -> interrupts
    ])
    cues = music.build_cues(tl, {"A5"}, gap_seconds=1.0)
    assert [c.song for c in cues] == ["song1", "song1", "song2"]
    assert (cues[0].rec_in, cues[0].rec_out) == (0, 200)     # add-edit merged
    assert (cues[1].rec_in, cues[1].rec_out) == (400, 500)
    assert (cues[2].rec_in, cues[2].rec_out) == (505, 600)


def test_music_interrupt_beats_gap():
    """A different piece coming in within the gap window still splits the cue."""
    from timeline_reader import music
    tl = _music_timeline([
        _music_clip("A5", "song1", 0, 100),
        _music_clip("A5", "song2", 105, 200),  # 5f gap but a different song
    ])
    cues = music.build_cues(tl, {"A5"}, gap_seconds=1.0)
    assert [c.song for c in cues] == ["song1", "song2"]


def test_music_checkerboard_across_tracks():
    """One cue laid across two tracks (so it can overlap itself) stays one cue,
    and lists both tracks."""
    from timeline_reader import music
    tl = _music_timeline([
        _music_clip("A5", "song1", 0, 100),
        _music_clip("A6", "song1", 90, 220),
    ])
    cues = music.build_cues(tl, {"A5", "A6"}, gap_seconds=1.0)
    assert len(cues) == 1
    assert cues[0].tracks == ["A5", "A6"]
    assert (cues[0].rec_in, cues[0].rec_out) == (0, 220)


def test_music_track_filtering():
    from timeline_reader import music
    tl = _music_timeline([
        _music_clip("A5", "song1", 0, 100),
        _music_clip("A1", "dialogue", 0, 100),  # not a chosen music track
    ])
    cues = music.build_cues(tl, {"A5"}, gap_seconds=1.0)
    assert len(cues) == 1 and cues[0].song == "song1"


def test_music_reel_from_hour():
    from timeline_reader import music
    from timeline_reader.timecode import frames_to_tc
    start = Timecode.from_string("01:00:00:00", 25.0).frames  # reel 1
    tl = _music_timeline([_music_clip("A5", "song1", 4592, 5000, artist="Anil")],
                         start_tc=start)
    cue = music.build_cues(tl, {"A5"})[0]
    assert cue.reel() == "1"
    assert frames_to_tc(cue.rec_in, cue.fps) == "01:03:03:17"  # start + rec_start
    assert cue.artist == "Anil"


def test_music_columns_track_is_song_title_and_album():
    """The Track column is the song name (title metadata), Filename is the file,
    and an Album column carries the album."""
    from timeline_reader import music
    from timeline_reader.columns import ColumnSelection
    tl = _music_timeline([
        _music_clip("A5", "01-Kashmiri", 0, 100, artist="Anil",
                    title="Kashmiri", album="Kashmiri Dhol"),
    ])
    sel = ColumnSelection(music.REPORT_KEY)  # defaults
    cues, cols, rows = music.build_rows(tl, {"A5"}, sel)
    headers = [c.label for c in cols]
    assert headers == ["Reel", "Track", "Artist / Composer", "Album",
                       "Filename", "TC In", "TC Out", "Duration"]
    row = dict(zip(headers, rows[0]))
    assert row["Track"] == "Kashmiri"            # song name, not the audio track
    assert row["Album"] == "Kashmiri Dhol"
    assert row["Filename"] == "01-Kashmiri.wav"  # the media file
    assert row["Artist / Composer"] == "Anil"


def test_music_muted_excluded_by_default():
    from timeline_reader import music
    from timeline_reader.columns import ColumnSelection
    tl = _music_timeline([
        _music_clip("A5", "song1", 0, 100),
        _music_clip("A6", "song2", 200, 300, muted=True),  # on a muted track
    ])
    # Default: muted clips dropped, no Muted column.
    cues = music.build_cues(tl, {"A5", "A6"})
    assert [c.song for c in cues] == ["song1"]

    # Included: the muted cue appears and is flagged in a Muted column.
    sel = ColumnSelection(music.REPORT_KEY)
    cues, cols, rows = music.build_rows(tl, {"A5", "A6"}, sel, include_muted=True)
    assert "Muted" in [c.label for c in cols]
    by_song = {c.song: c for c in cues}
    assert by_song["song2"].muted is True and by_song["song1"].muted is False
    muted_col = [c.label for c in cols].index("Muted")
    song2_row = rows[[c.song for c in cues].index("song2")]
    assert song2_row[muted_col] == "Muted"


def test_music_dissolve_inclusion():
    from timeline_reader import music
    clip = _music_clip("A5", "song1", 100, 200, head=25, tail=25)
    incl = music.build_cues(_music_timeline([clip]), {"A5"}, include_dissolves=True)[0]
    excl = music.build_cues(_music_timeline([clip]), {"A5"}, include_dissolves=False)[0]
    assert (incl.rec_in, incl.rec_out) == (100, 225)   # fades counted
    assert (excl.rec_in, excl.rec_out) == (125, 200)   # trimmed to hard cuts


def test_music_from_sample_bin():
    """End-to-end against the real sample bin's music tracks (A15/A16), if present."""
    path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "01-test files", "bin", "04-Lock.avb",
    )
    if not os.path.exists(path):
        return
    from timeline_reader import music
    from timeline_reader.parsers import parse_timeline

    tl = parse_timeline(path)
    assert tl.audio_clips, "sound tracks should be parsed into audio_clips"
    assert {"A15", "A16"} <= set(music.available_tracks(tl))

    cues = music.build_cues(tl, {"A15", "A16"}, gap_seconds=1.0)
    # The Kashmiri cue opens the reel and is checkerboarded across A15 + A16.
    first = cues[0]
    assert first.song == "01-Kashmiri"
    assert first.title == "Kashmiri"          # Track column = song name
    assert first.album == "Kashmiri Dhol"
    assert first.tracks == ["A15", "A16"]      # kept for the optional Audio Track col
    assert first.artist == "Anil"
    assert first.muted is False                # nothing muted in the sample bin
    # A song reused far later in the timeline splits into separate cues.
    dimensions = [c for c in cues if c.song == "13-Dimensions"]
    assert len(dimensions) >= 2


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all tests passed")
